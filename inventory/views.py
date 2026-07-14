import io
import os
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.contrib import messages
from django.contrib.contenttypes.models import ContentType
from django.db.models import Sum, F, DecimalField, ExpressionWrapper
from django.http import HttpResponse, HttpResponseRedirect, FileResponse
from django.urls import reverse
from django.conf import settings

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from .models import Supplier, InventoryItem, PurchaseOrder, POLineItem, SupplierCatalogItem
from .forms import InventoryItemForm, SupplierForm, PurchaseOrderForm, SupplierCatalogProductForm
from core.models import AuditLog, State, Transition, Organization
from core.reports.graphs import GraphGenerator

@login_required
def dashboard(request):
    """
    Overview page with key KPIs and interactive charts.
    """
    total_items = InventoryItem.objects.count()
    low_stock_items = InventoryItem.objects.filter(stock_level__lte=F('reorder_threshold')).count()
    total_suppliers = Supplier.objects.count()
    
    # Calculate stock value: Sum of (stock_level * unit_price)
    stock_value = InventoryItem.objects.aggregate(
        total=Sum(ExpressionWrapper(F('stock_level') * F('unit_price'), output_field=DecimalField()))
    )['total'] or 0.00
    
    # Generate Chart 1: Stock levels bar chart
    items_qs = InventoryItem.objects.all()[:15]
    stock_chart = GraphGenerator.generate_bar_chart(
        queryset=items_qs,
        x_field='name',
        y_field='stock_level',
        title="Stock Levels by Product"
    )
    
    # Generate Chart 2: Purchase Order status distribution
    po_states_data = list(
        PurchaseOrder.objects.values('workflow_state__name')
        .annotate(total=Sum('total_amount'))
    )
    # Rename key for the pie chart generator
    for item in po_states_data:
        item['status_name'] = item['workflow_state__name'] or 'Draft'
        
    po_chart = GraphGenerator.generate_pie_chart(
        queryset=po_states_data,
        names_field='status_name',
        values_field='total',
        title="Purchase Orders by Status"
    )

    # Generate Chart 3: Demand Forecast Chart based on received PO line item history
    from django.db.models.functions import TruncDate
    historical_orders = list(
        POLineItem.objects.filter(purchase_order__workflow_state__name='Received')
        .annotate(order_date=TruncDate('purchase_order__created_at'))
        .values('order_date')
        .annotate(quantity_ordered=Sum('quantity'))
        .order_by('order_date')
    )
    
    # Format dates and quantities for Plotly
    plotly_data = []
    for order in historical_orders:
        plotly_data.append({
            'date': str(order['order_date']),
            'total_qty': float(order['quantity_ordered'])
        })
        
    forecast_chart = GraphGenerator.generate_forecast_chart(
        historical_data=plotly_data,
        x_field='date',
        y_field='total_qty',
        title="Procurement Demand Trend & Forecast"
    )

    context = {
        'total_items': total_items,
        'low_stock_items': low_stock_items,
        'total_suppliers': total_suppliers,
        'stock_value': stock_value,
        'stock_chart': stock_chart,
        'po_chart': po_chart,
        'forecast_chart': forecast_chart
    }
    return render(request, 'inventory/dashboard.html', context)

@login_required
def item_list(request):
    """List of all inventory items with filter options."""
    query = request.GET.get('q', '')
    low_stock = request.GET.get('low_stock', '')

    items = InventoryItem.objects.all().prefetch_related('catalog_items__supplier').order_by('sku')
    if query:
        items = items.filter(name__icontains=query) | items.filter(sku__icontains=query)
    if low_stock == '1':
        items = items.filter(stock_level__lte=F('reorder_threshold'))

    # Pagination
    from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
    paginator = Paginator(items, 10)
    page = request.GET.get('page')
    try:
        paginated_items = paginator.page(page)
    except PageNotAnInteger:
        paginated_items = paginator.page(1)
    except EmptyPage:
        paginated_items = paginator.page(paginator.num_pages)

    # Build query string excluding 'page'
    params = request.GET.copy()
    if 'page' in params:
        del params['page']
    query_string = params.urlencode()

    context = {
        'items': paginated_items,
        'query': query,
        'low_stock': low_stock,
        'query_string': query_string
    }
    return render(request, 'inventory/item_list.html', context)

@login_required
def supplier_list(request):
    """List of all suppliers."""
    suppliers_list = Supplier.objects.all().prefetch_related('supplied_items').order_by('name')
    
    # Pagination
    from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
    paginator = Paginator(suppliers_list, 10)
    page = request.GET.get('page')
    try:
        paginated_suppliers = paginator.page(page)
    except PageNotAnInteger:
        paginated_suppliers = paginator.page(1)
    except EmptyPage:
        paginated_suppliers = paginator.page(paginator.num_pages)
        
    # Build query string excluding 'page'
    params = request.GET.copy()
    if 'page' in params:
        del params['page']
    query_string = params.urlencode()
    
    context = {
        'suppliers': paginated_suppliers,
        'query_string': query_string
    }
    return render(request, 'inventory/supplier_list.html', context)

@login_required
def supplier_detail(request, pk):
    """Detailed view of a supplier, showing basic info and product catalog."""
    supplier = get_object_or_404(Supplier, pk=pk)
    catalog_items = supplier.catalog_entries.all().select_related('item')
    return render(request, 'inventory/supplier_detail.html', {
        'supplier': supplier,
        'catalog_items': catalog_items
    })

@login_required
def supplier_toggle_status(request, pk):
    """Toggle the active/inactive status of a supplier and update associated products."""
    from django.contrib import messages
    from django.views.decorators.http import require_POST
    
    supplier = get_object_or_404(Supplier, pk=pk)
    if request.method == 'POST':
        supplier.is_active = not supplier.is_active
        supplier.save()
        
        # Update all items supplied by this supplier
        supplier.supplied_items.all().update(is_active=supplier.is_active)
        
        status_str = "activated" if supplier.is_active else "deactivated"
        messages.success(request, f"Supplier '{supplier.name}' was successfully {status_str}.")
    
    return redirect(reverse('inventory:supplier_detail', args=[supplier.id]))

def get_filtered_pos(request):
    supplier_id = request.GET.get('supplier')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    
    from datetime import datetime, timedelta
    import calendar
    from django.utils import timezone
    
    start_dt = None
    end_dt = None
    error_message = None
    
    # Parse dates
    if start_date:
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d").date()
        except ValueError:
            error_message = "Invalid start date format. Use YYYY-MM-DD."
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d").date()
        except ValueError:
            error_message = "Invalid end date format. Use YYYY-MM-DD."
            
    # If no dates are specified, default to the current month (first to last day)
    if not start_date and not end_date:
        today = timezone.now().date()
        start_dt = today.replace(day=1)
        last_day = calendar.monthrange(today.year, today.month)[1]
        end_dt = today.replace(day=last_day)
        start_date = start_dt.strftime("%Y-%m-%d")
        end_date = end_dt.strftime("%Y-%m-%d")
    
    # Validate date range requirements
    if not error_message:
        if start_dt and end_dt:
            if (end_dt - start_dt).days > 31:
                error_message = "Date range cannot exceed 31 days. Please select a shorter range."
            elif (end_dt - start_dt).days < 0:
                error_message = "Start date must be before or equal to end date."
        else:
            # If one is provided and the other is missing
            error_message = "Both Start Date and End Date must be selected."
            
    pos = PurchaseOrder.objects.all().select_related('supplier', 'workflow_state').order_by('-created_at')
    
    if supplier_id:
        pos = pos.filter(supplier_id=supplier_id)
        
    if not error_message and start_dt and end_dt:
        pos = pos.filter(created_at__date__range=[start_dt, end_dt])
    else:
        # If there's an error, don't return any PO records to avoid slow database scans
        pos = pos.none()
        
    return pos, supplier_id, start_date, end_date, error_message

@login_required
def po_list(request):
    """List of all purchase orders with filter options."""
    pos, supplier_id, start_date, end_date, error_message = get_filtered_pos(request)
    if error_message:
        from django.contrib import messages
        messages.error(request, error_message)
        
    suppliers = Supplier.objects.all()
    
    # Pagination
    from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
    paginator = Paginator(pos, 10)
    page = request.GET.get('page')
    try:
        paginated_pos = paginator.page(page)
    except PageNotAnInteger:
        paginated_pos = paginator.page(1)
    except EmptyPage:
        paginated_pos = paginator.page(paginator.num_pages)
        
    # Construct query string for export links
    import urllib.parse
    params = request.GET.copy()
    for k in list(params.keys()):
        if not params[k]:
            del params[k]
    # Set dates in the params so the export links have them even if the user didn't enter them (defaulting to current month)
    if start_date and 'start_date' not in params:
        params['start_date'] = start_date
    if end_date and 'end_date' not in params:
        params['end_date'] = end_date
        
    # Exclude page from query_string to avoid appending page number twice in pagination links
    pagination_params = params.copy()
    if 'page' in pagination_params:
        del pagination_params['page']
    query_string = pagination_params.urlencode()
    
    context = {
        'purchase_orders': paginated_pos,
        'suppliers': suppliers,
        'selected_supplier_id': int(supplier_id) if supplier_id else None,
        'start_date': start_date,
        'end_date': end_date,
        'query_string': query_string
    }
    return render(request, 'inventory/po_list.html', context)

def get_organization_header_flowables(styles):
    """
    Returns a list of flowables (header table, divider line, spacer) containing
    organization details and logo to be included at the top of PDF reports.
    """
    from reportlab.platypus import Image as RLImage
    
    org = Organization.objects.first()
    org_name = org.name if org else "Quantum Global"
    org_address = org.address if org else "123 Business Road, Corporate Hub, India"
    org_email = org.email if org else "info@quantumglobal.com"
    org_phone = org.phone if org else "+91 99999 88888"
    org_gstin = org.gstin if org else "27AAAAA0000A1Z5"
    org_license = org.license_number if org else "LIC-998877"

    logo_path = None
    if org and org.logo and os.path.exists(org.logo.path):
        logo_path = org.logo.path
    if not logo_path:
        default_logo_path = os.path.join(settings.MEDIA_ROOT, 'org_logos', 'default_logo.png')
        if os.path.exists(default_logo_path):
            logo_path = default_logo_path

    org_details_html = f"""
    <b>{org_name}</b><br/>
    {org_address}<br/>
    Phone: {org_phone} | Email: {org_email}<br/>
    GSTIN: {org_gstin} | License: {org_license}
    """
    
    org_details_style = ParagraphStyle(
        'OrgDetails',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569')
    )
    
    org_paragraph = Paragraph(org_details_html, org_details_style)
    
    if logo_path:
        try:
            logo_flowable = RLImage(logo_path, width=50, height=50)
            header_table_data = [[logo_flowable, org_paragraph]]
            header_table = Table(header_table_data, colWidths=[60, 480])
        except Exception:
            header_table_data = [[org_paragraph]]
            header_table = Table(header_table_data, colWidths=[540])
    else:
        header_table_data = [[org_paragraph]]
        header_table = Table(header_table_data, colWidths=[540])

    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
        ('TOPPADDING', (0,0), (-1,-1), 0),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    
    divider = Table([['']], colWidths=[540], rowHeights=[1])
    divider.setStyle(TableStyle([
        ('LINEABOVE', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
    ]))
    
    return [header_table, divider, Spacer(1, 10)]

@login_required
def po_export_pdf(request):
    """View to download filtered purchase orders in PDF format."""
    from django.utils import timezone
    pos, supplier_id, start_date, end_date, error_message = get_filtered_pos(request)
    
    if error_message:
        from django.contrib import messages
        messages.error(request, error_message)
        # Redirect back to list page preserving current filter params
        import urllib.parse
        q_params = urllib.parse.urlencode({
            'supplier': supplier_id or '',
            'start_date': start_date or '',
            'end_date': end_date or ''
        })
        return redirect(reverse('inventory:po_list') + '?' + q_params)
        
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    story = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e3a8a'),
        spaceAfter=8
    )
    subtitle_style = ParagraphStyle(
        'ReportSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#475569'),
        spaceAfter=12
    )
    po_header_style = ParagraphStyle(
        'POHeader',
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=8,
        spaceAfter=3
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1e293b')
    )
    body_bold = ParagraphStyle(
        'DocBodyBold',
        parent=body_style,
        fontName='Helvetica-Bold'
    )
    header_style = ParagraphStyle(
        'HeaderStyle',
        parent=body_style,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )
    
    story.extend(get_organization_header_flowables(styles))
    story.append(Paragraph("Purchase Orders Report", title_style))
    
    # Filter text
    filter_parts = []
    if supplier_id:
        supplier_name = Supplier.objects.filter(id=supplier_id).values_list('name', flat=True).first()
        filter_parts.append(f"Supplier: {supplier_name or 'All'}")
    if start_date:
        filter_parts.append(f"From: {start_date}")
    if end_date:
        filter_parts.append(f"To: {end_date}")
    filter_text = ", ".join(filter_parts) if filter_parts else "Filters: None"
    story.append(Paragraph(f"Generated on {timezone.now().strftime('%Y-%m-%d %H:%M')} | {filter_text}", subtitle_style))
    
    for po in pos:
        story.append(Spacer(1, 8))
        po_info = f"<b>PO Number:</b> {po.po_number} | <b>Supplier:</b> {po.supplier.name} | <b>Date:</b> {po.created_at.strftime('%Y-%m-%d')} | <b>Status:</b> {po.workflow_state.name if po.workflow_state else 'Draft'} | <b>Total:</b> Rs. {po.total_amount:.2f}"
        story.append(Paragraph(po_info, po_header_style))
        
        headers = [
            Paragraph("SKU", header_style),
            Paragraph("Product Name", header_style),
            Paragraph("Quantity", header_style),
            Paragraph("Unit Price", header_style),
            Paragraph("Subtotal", header_style)
        ]
        table_data = [headers]
        for line in po.lines.all().select_related('item'):
            table_data.append([
                Paragraph(line.item.sku, body_style),
                Paragraph(line.item.name, body_style),
                Paragraph(str(line.quantity), body_style),
                Paragraph(f"Rs. {line.unit_price:.2f}", body_style),
                Paragraph(f"Rs. {line.subtotal:.2f}", body_style)
            ])
            
        t = Table(table_data, colWidths=[100, 200, 60, 80, 100])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#475569')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ]))
        story.append(t)
        story.append(Spacer(1, 3))
        
    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="Purchase_Orders_Report.pdf"'
    return response

def inject_excel_organization_header(ws):
    """
    Injects organization logo and details at the top of an openpyxl worksheet.
    Returns the next available row index for content.
    """
    from core.models import Organization
    from django.conf import settings
    from openpyxl.styles import Font, Border, Side
    from openpyxl.drawing.image import Image as OpenpyxlImage
    import os

    org = Organization.objects.first()
    org_name = org.name if org else "Quantum Global"
    org_address = org.address if org else "123 Business Road, Corporate Hub, India"
    org_email = org.email if org else "info@quantumglobal.com"
    org_phone = org.phone if org else "+91 99999 88888"
    org_gstin = org.gstin if org else "27AAAAA0000A1Z5"
    org_license = org.license_number if org else "LIC-998877"

    logo_path = None
    if org and org.logo and os.path.exists(org.logo.path):
        logo_path = org.logo.path
    if not logo_path:
        default_logo_path = os.path.join(settings.MEDIA_ROOT, 'org_logos', 'default_logo.png')
        if os.path.exists(default_logo_path):
            logo_path = default_logo_path

    # Define fonts
    org_name_font = Font(name='Arial', size=14, bold=True, color='1E3A8A')
    org_details_font = Font(name='Arial', size=9, color='475569')

    # Add details in Column B
    ws.cell(row=1, column=2, value=org_name).font = org_name_font
    ws.cell(row=2, column=2, value=org_address).font = org_details_font
    ws.cell(row=3, column=2, value=f"Phone: {org_phone} | Email: {org_email}").font = org_details_font
    ws.cell(row=4, column=2, value=f"GSTIN: {org_gstin} | License: {org_license}").font = org_details_font

    if logo_path:
        try:
            img = OpenpyxlImage(logo_path)
            img.width = 55
            img.height = 55
            ws.add_image(img, 'A1')
        except Exception:
            pass

    # Row heights
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 14
    ws.row_dimensions[3].height = 14
    ws.row_dimensions[4].height = 14
    ws.row_dimensions[5].height = 10

    # Draw bottom border on cells in row 5
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    bottom_border = Border(bottom=thin_border_side)
    for col in range(1, 11):
        ws.cell(row=5, column=col).border = bottom_border

    return 7


@login_required
def po_export_excel(request):
    """View to download filtered purchase orders in Excel format."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    pos, supplier_id, start_date, end_date, error_message = get_filtered_pos(request)
    
    if error_message:
        from django.contrib import messages
        messages.error(request, error_message)
        import urllib.parse
        q_params = urllib.parse.urlencode({
            'supplier': supplier_id or '',
            'start_date': start_date or '',
            'end_date': end_date or ''
        })
        return redirect(reverse('inventory:po_list') + '?' + q_params)
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Orders"
    
    # Inject organization details and logo at the top
    start_row = inject_excel_organization_header(ws)
    
    title_font = Font(name='Arial', size=16, bold=True, color='1E3A8A')
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    bold_font = Font(name='Arial', size=10, bold=True)
    normal_font = Font(name='Arial', size=10)
    
    # Title
    ws.cell(row=start_row, column=1, value="Purchase Orders Report").font = title_font
    
    # Filter descriptions
    ws.cell(row=start_row + 2, column=1, value="Filters Applied:").font = bold_font
    filter_row = start_row + 2
    if supplier_id:
        supplier_name = Supplier.objects.filter(id=supplier_id).values_list('name', flat=True).first()
        ws.cell(row=filter_row, column=2, value=f"Supplier: {supplier_name}").font = normal_font
        filter_row += 1
    if start_date:
        ws.cell(row=filter_row, column=2, value=f"Start Date: {start_date}").font = normal_font
        filter_row += 1
    if end_date:
        ws.cell(row=filter_row, column=2, value=f"End Date: {end_date}").font = normal_font
        filter_row += 1
    if filter_row == start_row + 2:
        ws.cell(row=start_row + 2, column=2, value="None").font = normal_font
        filter_row += 1
        
    headers = [
        "PO Number", "Supplier", "Date Created", "Status", 
        "Item SKU", "Item Name", "Quantity", "Unit Cost (₹)", "Line Subtotal (₹)", "PO Total Amount (₹)"
    ]
    
    start_data_row = filter_row + 2
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_data_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        
    row_idx = start_data_row + 1
    for po in pos:
        po_num = po.po_number
        sup_name = po.supplier.name
        po_date = po.created_at.strftime('%Y-%m-%d %H:%M')
        po_status = po.workflow_state.name if po.workflow_state else 'Draft'
        po_total = float(po.total_amount)
        
        lines = po.lines.all().select_related('item')
        if not lines.exists():
            ws.cell(row=row_idx, column=1, value=po_num).font = normal_font
            ws.cell(row=row_idx, column=2, value=sup_name).font = normal_font
            ws.cell(row=row_idx, column=3, value=po_date).font = normal_font
            ws.cell(row=row_idx, column=4, value=po_status).font = normal_font
            ws.cell(row=row_idx, column=10, value=po_total).font = bold_font
            row_idx += 1
        else:
            for line in lines:
                ws.cell(row=row_idx, column=1, value=po_num).font = normal_font
                ws.cell(row=row_idx, column=2, value=sup_name).font = normal_font
                ws.cell(row=row_idx, column=3, value=po_date).font = normal_font
                ws.cell(row=row_idx, column=4, value=po_status).font = normal_font
                ws.cell(row=row_idx, column=5, value=line.item.sku).font = normal_font
                ws.cell(row=row_idx, column=6, value=line.item.name).font = normal_font
                ws.cell(row=row_idx, column=7, value=line.quantity).font = normal_font
                ws.cell(row=row_idx, column=8, value=float(line.unit_price)).font = normal_font
                ws.cell(row=row_idx, column=9, value=float(line.subtotal)).font = normal_font
                ws.cell(row=row_idx, column=10, value=po_total).font = bold_font
                row_idx += 1
                
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Purchase_Orders_Report.xlsx"'
    wb.save(response)
    return response

@login_required
def po_detail(request, pk):
    """Detailed view of a Purchase Order, including its lines, workflow actions, and audit logs."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    
    # Fetch audit logs
    content_type = ContentType.objects.get_for_model(po)
    audit_logs = AuditLog.objects.filter(
        content_type=content_type,
        object_id=po.id
    ).select_related('user').order_by('-timestamp')

    context = {
        'instance': po,
        'lines': po.lines.all().select_related('item'),
        'transitions': po.get_available_transitions(request.user),
        'app_label': 'inventory',
        'model_name': 'purchaseorder',
        'audit_logs': audit_logs
    }
    return render(request, 'inventory/po_detail.html', context)

@login_required
@transaction.atomic
def po_create(request):
    """Create a new Purchase Order along with its lines."""
    suppliers = Supplier.objects.filter(is_active=True)
    items = InventoryItem.objects.filter(is_active=True).prefetch_related('catalog_items')

    if request.method == 'POST':
        po_number = request.POST.get('po_number')
        supplier_id = request.POST.get('supplier')
        
        supplier = get_object_or_404(Supplier, id=supplier_id)
        
        # Create PurchaseOrder (starts in the Initial state 'Draft')
        draft_state = State.objects.filter(workflow__model_name='inventory.PurchaseOrder', is_initial=True).first()
        
        po_kwargs = {
            'supplier': supplier,
            'workflow_state': draft_state,
        }
        if po_number:
            po_kwargs['po_number'] = po_number
            
        po = PurchaseOrder.objects.create(**po_kwargs)
        
        # Read form items
        item_ids = request.POST.getlist('item[]')
        quantities = request.POST.getlist('quantity[]')
        prices = request.POST.getlist('price[]')
        
        total = 0.00
        for item_id, qty, price in zip(item_ids, quantities, prices):
            if not item_id or not qty or not price:
                continue
            item = get_object_or_404(InventoryItem, id=item_id)
            q = int(qty)
            p = float(price)
            POLineItem.objects.create(
                purchase_order=po,
                item=item,
                quantity=q,
                unit_price=p
            )
            total += (q * p)
            
        po.total_amount = total
        po.save()
        
        return redirect(reverse('inventory:po_detail', args=[po.id]))

    context = {
        'suppliers': suppliers,
        'items': items
    }
    return render(request, 'inventory/po_create.html', context)

def generate_po_pdf_bytes(po):
    """Generates a professional PDF receipt for a Purchase Order and returns bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    story = []
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#1e3a8a'), # Premium Navy
        spaceAfter=15
    )
    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#334155'), # Slate Gray
        spaceBefore=10,
        spaceAfter=10
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1e293b')
    )
    body_bold = ParagraphStyle(
        'DocBodyBold',
        parent=body_style,
        fontName='Helvetica-Bold'
    )
    header_style = ParagraphStyle(
        'HeaderStyle',
        parent=body_style,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )

    # 1. Organization Header & Document Title
    story.extend(get_organization_header_flowables(styles))
    story.append(Paragraph("PURCHASE ORDER", title_style))
    story.append(Spacer(1, 10))

    # 2. Metadata details (2-column layout)
    meta_data = [
        [
            Paragraph(f"<b>PO Number:</b> {po.po_number}", body_style),
            Paragraph(f"<b>Supplier:</b> {po.supplier.name}", body_style)
        ],
        [
            Paragraph(f"<b>Date Created:</b> {po.created_at.strftime('%Y-%m-%d %H:%M')}", body_style),
            Paragraph(f"<b>Contact Email:</b> {po.supplier.contact_email}", body_style)
        ],
        [
            Paragraph(f"<b>Status:</b> {po.workflow_state.name if po.workflow_state else 'Draft'}", body_style),
            Paragraph(f"<b>Phone:</b> {po.supplier.phone or 'N/A'}", body_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 20))

    # 3. Line Items Table
    story.append(Paragraph("Line Items", section_heading))
    
    # Table headers
    headers = [
        Paragraph("SKU", header_style),
        Paragraph("Product Name", header_style),
        Paragraph("Quantity", header_style),
        Paragraph("Unit Price", header_style),
        Paragraph("Subtotal", header_style)
    ]
    table_data = [headers]
    
    # Table rows
    for line in po.lines.all().select_related('item'):
        table_data.append([
            Paragraph(line.item.sku, body_style),
            Paragraph(line.item.name, body_style),
            Paragraph(str(line.quantity), body_style),
            Paragraph(f"Rs. {line.unit_price:.2f}", body_style),
            Paragraph(f"Rs. {line.subtotal:.2f}", body_bold)
        ])
        
    # Total row
    table_data.append([
        Paragraph("<b>Total Amount:</b>", body_bold),
        "", "", "",
        Paragraph(f"<b>Rs. {po.total_amount:.2f}</b>", body_bold)
    ])
    
    items_table = Table(table_data, colWidths=[100, 200, 70, 80, 90])
    items_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e3a8a')),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('GRID', (0,0), (-1,-2), 0.5, colors.HexColor('#e2e8f0')),
        ('SPAN', (0,-1), (3,-1)), # Span "Total Amount" label across first 4 columns
        ('ALIGN', (0,-1), (3,-1), 'RIGHT'),
        ('TOPPADDING', (0,-1), (-1,-1), 10),
        ('LINEABOVE', (0,-1), (-1,-1), 1, colors.HexColor('#94a3b8')),
    ]))
    story.append(items_table)
    
    # Build Document
    doc.build(story)
    
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


@login_required
def po_download_pdf(request, pk):
    """Generates and downloads a professional PDF receipt for a Purchase Order."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    pdf_bytes = generate_po_pdf_bytes(po)
    
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="PO_{po.po_number}.pdf"'
    return response

# --- InventoryItem CRUD ---
@login_required
def item_create(request):
    if request.method == 'POST':
        form = InventoryItemForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('inventory:item_list')
    else:
        form = InventoryItemForm()
    return render(request, 'inventory/form.html', {'form': form, 'title': 'Create Product', 'back_url': reverse('inventory:item_list')})

@login_required
def item_update(request, pk):
    item = get_object_or_404(InventoryItem, pk=pk)
    if request.method == 'POST':
        form = InventoryItemForm(request.POST, instance=item)
        if form.is_valid():
            form.save()
            return redirect('inventory:item_list')
    else:
        form = InventoryItemForm(instance=item)
    return render(request, 'inventory/form.html', {'form': form, 'title': f'Edit Product: {item.sku}', 'back_url': reverse('inventory:item_list')})

@login_required
def item_delete(request, pk):
    item = get_object_or_404(InventoryItem, pk=pk)
    if request.method == 'POST':
        item.delete()
        return redirect('inventory:item_list')
    return render(request, 'inventory/delete_confirm.html', {'object': item, 'title': f'Delete Product: {item.sku}', 'back_url': reverse('inventory:item_list')})


# --- Supplier CRUD ---
@login_required
def supplier_create(request):
    if request.method == 'POST':
        form = SupplierForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('inventory:supplier_list')
    else:
        form = SupplierForm()
    return render(request, 'inventory/form.html', {'form': form, 'title': 'Create Supplier', 'back_url': reverse('inventory:supplier_list')})

@login_required
def supplier_update(request, pk):
    supplier = get_object_or_404(Supplier, pk=pk)
    if request.method == 'POST':
        form = SupplierForm(request.POST, instance=supplier)
        if form.is_valid():
            form.save()
            return redirect('inventory:supplier_list')
    else:
        form = SupplierForm(instance=supplier)
    return render(request, 'inventory/form.html', {'form': form, 'title': f'Edit Supplier: {supplier.name}', 'back_url': reverse('inventory:supplier_list')})

@login_required
def supplier_delete(request, pk):
    supplier = get_object_or_404(Supplier, pk=pk)
    if request.method == 'POST':
        supplier.delete()
        return redirect('inventory:supplier_list')
    return render(request, 'inventory/delete_confirm.html', {'object': supplier, 'title': f'Delete Supplier: {supplier.name}', 'back_url': reverse('inventory:supplier_list')})

@login_required
def supplier_catalog_modal(request, pk, form=None):
    """Render the catalog management modal for a supplier."""
    supplier = get_object_or_404(Supplier, pk=pk)
    items = InventoryItem.objects.all().order_by('name')
    catalog_items = supplier.catalog_entries.all().select_related('item')
    if form is None:
        form = SupplierCatalogProductForm()
    context = {
        'supplier': supplier,
        'items': items,
        'catalog_items': catalog_items,
        'form': form,
    }
    return render(request, 'inventory/supplier_catalog_modal.html', context)

@login_required
@transaction.atomic
def supplier_catalog_add_product(request, pk):
    """Create a new product and add it to the supplier's catalog."""
    supplier = get_object_or_404(Supplier, pk=pk)
    if request.method == 'POST':
        form = SupplierCatalogProductForm(request.POST)
        if form.is_valid():
            item = form.save()
            SupplierCatalogItem.objects.create(
                supplier=supplier,
                item=item,
                price=form.cleaned_data['supplier_price']
            )
            return supplier_catalog_modal(request, pk)
    else:
        form = SupplierCatalogProductForm()
    return supplier_catalog_modal(request, pk, form=form)

@login_required
@transaction.atomic
def supplier_catalog_remove(request, entry_pk):
    """Remove an item from the supplier's catalog."""
    entry = get_object_or_404(SupplierCatalogItem, pk=entry_pk)
    supplier_id = entry.supplier_id
    entry.delete()
    return supplier_catalog_modal(request, supplier_id)


# --- PurchaseOrder Delete ---
@login_required
def po_delete(request, pk):
    po = get_object_or_404(PurchaseOrder, pk=pk)
    if po.workflow_state and po.workflow_state.name != 'Draft':
        return HttpResponse("Deletion only allowed for Draft Purchase Orders.", status=403)
        
    if request.method == 'POST':
        po.delete()
        return redirect('inventory:po_list')
    return render(request, 'inventory/delete_confirm.html', {'object': po, 'title': f'Delete Purchase Order: {po.po_number}', 'back_url': reverse('inventory:po_detail', args=[po.id])})


@login_required
def item_price_history(request, pk):
    """
    Renders price history for a specific inventory product.
    If requested via HTMX, returns an overlay modal fragment.
    """
    item = get_object_or_404(InventoryItem, pk=pk)
    logs = list(item.price_logs.all().order_by('changed_at'))
    
    price_logs_with_diff = []
    prev_price = None
    
    for log in logs:
        diff_pct = 0.0
        if prev_price is not None and prev_price != 0:
            diff_pct = float((log.price - prev_price) / prev_price) * 100
        
        log.diff_pct = diff_pct
        prev_price = log.price
        price_logs_with_diff.append(log)
        
    price_logs_with_diff.reverse()
    
    context = {
        'item': item,
        'price_logs': price_logs_with_diff,
    }
    
    if request.headers.get('HX-Request'):
        return render(request, 'inventory/item_price_history_modal.html', context)
        
    return render(request, 'inventory/item_price_history.html', context)


@login_required
@transaction.atomic
def po_edit(request, pk):
    """Edit a Purchase Order in Draft state."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    
    # Security/state check: Only Draft POs can be edited
    if po.workflow_state and po.workflow_state.name != 'Draft':
        return HttpResponse("Editing is only allowed for Draft Purchase Orders.", status=403)
        
    suppliers = Supplier.objects.filter(is_active=True) | Supplier.objects.filter(id=po.supplier_id)
    items = InventoryItem.objects.filter(is_active=True).prefetch_related('catalog_items')
    
    if request.method == 'POST':
        supplier_id = request.POST.get('supplier')
        supplier = get_object_or_404(Supplier, id=supplier_id)
        
        po.supplier = supplier
        
        # Clear existing line items and rebuild
        po.lines.all().delete()
        
        # Read form items
        item_ids = request.POST.getlist('item[]')
        quantities = request.POST.getlist('quantity[]')
        prices = request.POST.getlist('price[]')
        
        total = 0.00
        for item_id, qty, price in zip(item_ids, quantities, prices):
            if not item_id or not qty or not price:
                continue
            item = get_object_or_404(InventoryItem, id=item_id)
            q = int(qty)
            p = float(price)
            POLineItem.objects.create(
                purchase_order=po,
                item=item,
                quantity=q,
                unit_price=p
            )
            total += (q * p)
            
        po.total_amount = total
        po.save()
        
        return redirect(reverse('inventory:po_detail', args=[po.id]))
        
    # Serialize existing line items to JSON for Alpine.js initialization
    import json
    lines_data = []
    for line in po.lines.all():
        lines_data.append({
            'item_id': str(line.item.id),
            'quantity': line.quantity,
            'price': float(line.unit_price)
        })
        
    context = {
        'po': po,
        'suppliers': suppliers,
        'items': items,
        'lines_json': json.dumps(lines_data),
    }
    return render(request, 'inventory/po_edit.html', context)


@login_required
def po_receive_modal(request, pk, transition_id):
    """Render the modal content for validating the received stock."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    transition = get_object_or_404(Transition, id=transition_id)
    return render(request, 'inventory/po_receive_modal.html', {
        'po': po,
        'transition': transition
    })


@login_required
@transaction.atomic
def po_receive_submit(request, pk, transition_id):
    """Process the submitted received stock validation and transition the PO state."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    transition = get_object_or_404(Transition, id=transition_id)
    
    # Update lines with actual received values and expiry dates
    line_ids = request.POST.getlist('line_id[]')
    
    for line_id in line_ids:
        line = get_object_or_404(POLineItem, id=line_id, purchase_order=po)
        
        qty_str = request.POST.get(f'received_qty_{line_id}')
        price_str = request.POST.get(f'received_price_{line_id}')
        expiry_date = request.POST.get(f'expiry_date_{line_id}')
        
        if qty_str is not None:
            line.received_quantity = int(qty_str)
        if price_str is not None:
            line.received_unit_price = float(price_str)
        if expiry_date:
            line.expiry_date = expiry_date
        line.save()
        
    # Execute workflow transition
    try:
        po.transition_to(transition, request.user)
    except ValueError as e:
        return HttpResponse(f"Transition error: {str(e)}", status=400)
        
    # Return HX-Refresh header to reload the full PO detail page with updated stock levels/totals
    response = HttpResponse()
    response['HX-Refresh'] = 'true'
    return response


from django.core.mail import EmailMessage

@login_required
def po_email_modal(request, pk):
    """Render the modal content for sending PO via email."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    
    if not po.workflow_state or po.workflow_state.name != "Approved":
        return HttpResponse("Email can only be sent for Approved Purchase Orders.", status=403)
        
    default_subject = f"Purchase Order {po.po_number} - ERP Framework"
    
    body_lines = [
        f"Dear {po.supplier.name},",
        "",
        f"Please find attached our Purchase Order {po.po_number} containing the following requested items:",
        ""
    ]
    for line in po.lines.all().select_related('item'):
        body_lines.append(f"- {line.item.name} (SKU: {line.item.sku}): {line.quantity} units @ Rs. {line.unit_price:.2f}/unit")
    body_lines.extend([
        "",
        f"Total Estimated Value: Rs. {po.total_amount:.2f}",
        "",
        "Please review the attached PDF document for our full terms and shipping/billing information.",
        "Confirm receipt and estimated delivery date by replying to this email.",
        "",
        "Best regards,",
        "Procurement Team"
    ])
    default_body = "\n".join(body_lines)
    
    return render(request, 'inventory/po_email_modal.html', {
        'po': po,
        'default_subject': default_subject,
        'default_body': default_body
    })


@login_required
def po_email_submit(request, pk):
    """Process the submitted PO email and open the local email client pre-filled with details."""
    po = get_object_or_404(PurchaseOrder, pk=pk)
    
    if not po.workflow_state or po.workflow_state.name != "Approved":
        return HttpResponse("Email can only be sent for Approved Purchase Orders.", status=403)
        
    recipient = request.POST.get('email')
    subject = request.POST.get('subject')
    body = request.POST.get('body')
    
    if not recipient or not subject or not body:
        return HttpResponse("All fields (Recipient, Subject, Body) are required.", status=400)
        
    from django.utils.html import escapejs
    
    # Return HTMX response with OOB toast and client-side JS redirection to mailto:
    response_content = (
        f'<div hx-swap-oob="beforeend:body">'
        f'  <div class="toast toast-end z-50">'
        f'    <div class="alert alert-success shadow-lg rounded-xl flex gap-2 font-semibold text-white">'
        f'      <span>Opening email client...</span>'
        f'    </div>'
        f'  </div>'
        f'</div>'
        f'<script>'
        f'  document.getElementById("po_email_modal")?.remove();'
        f'  const recipient = "{escapejs(recipient)}";'
        f'  const subject = "{escapejs(subject)}";'
        f'  const body = "{escapejs(body)}";'
        f'  window.location.href = "mailto:" + recipient + "?subject=" + encodeURIComponent(subject) + "&body=" + encodeURIComponent(body);'
        f'</script>'
    )
    return HttpResponse(response_content)


import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def generate_supplier_catalog_pdf_bytes(supplier, catalog_items):
    """Generates a professional PDF of the Supplier Catalog."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    story = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e3a8a'),
        spaceAfter=15
    )
    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#334155'),
        spaceBefore=12,
        spaceAfter=10
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#1e293b')
    )
    header_style = ParagraphStyle(
        'HeaderStyle',
        parent=body_style,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )
    
    story.extend(get_organization_header_flowables(styles))
    story.append(Paragraph(f"SUPPLIER PRODUCT CATALOG", title_style))
    story.append(Spacer(1, 10))
    
    meta_data = [
        [
            Paragraph(f"<b>Supplier Name:</b> {supplier.name}", body_style),
            Paragraph(f"<b>Contact Email:</b> {supplier.contact_email}", body_style)
        ],
        [
            Paragraph(f"<b>Phone:</b> {supplier.phone or 'N/A'}", body_style),
            Paragraph(f"<b>GST Number:</b> {supplier.gst_number or 'N/A'}", body_style)
        ],
        [
            Paragraph(f"<b>Address:</b> {supplier.address or 'N/A'}", body_style),
            Paragraph("", body_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 15))
    
    story.append(Paragraph("Catalog Items", section_heading))
    
    table_data = [[
        Paragraph("SKU", header_style),
        Paragraph("Product Name", header_style),
        Paragraph("Default Market Price", header_style),
        Paragraph("Supplier Cost", header_style),
        Paragraph("Expiry Req", header_style)
    ]]
    
    for entry in catalog_items:
        table_data.append([
            Paragraph(entry.item.sku, body_style),
            Paragraph(entry.item.name, body_style),
            Paragraph(f"Rs. {entry.item.unit_price:.2f}", body_style),
            Paragraph(f"<b>Rs. {entry.price:.2f}</b>", body_style),
            Paragraph("Yes" if entry.item.has_expiry_date else "No", body_style)
        ])
        
    catalog_table = Table(table_data, colWidths=[100, 200, 80, 80, 80])
    catalog_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e3a8a')),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    story.append(catalog_table)
    
    doc.build(story)
    return buffer.getvalue()

@login_required
def export_supplier_catalog_pdf(request, pk):
    """View to download the supplier catalog in PDF format."""
    supplier = get_object_or_404(Supplier, pk=pk)
    catalog_items = supplier.catalog_entries.all().select_related('item')
    pdf_bytes = generate_supplier_catalog_pdf_bytes(supplier, catalog_items)
    
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    filename = f"Catalog_{supplier.name.replace(' ', '_')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

@login_required
def export_supplier_catalog_excel(request, pk):
    """View to download the supplier catalog in Excel format (.xlsx)."""
    supplier = get_object_or_404(Supplier, pk=pk)
    catalog_items = supplier.catalog_entries.all().select_related('item')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Supplier Catalog"

    # Inject organization details and logo at the top
    start_row = inject_excel_organization_header(ws)

    title_font = Font(name='Arial', size=16, bold=True, color='1E3A8A')
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    bold_font = Font(name='Arial', size=10, bold=True)
    normal_font = Font(name='Arial', size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    ws.cell(row=start_row, column=1, value=f"Supplier Catalog - {supplier.name}").font = title_font
    
    ws.cell(row=start_row + 2, column=1, value="Email Address:").font = bold_font
    ws.cell(row=start_row + 2, column=2, value=supplier.contact_email).font = normal_font
    ws.cell(row=start_row + 3, column=1, value="Phone Number:").font = bold_font
    ws.cell(row=start_row + 3, column=2, value=supplier.phone or "N/A").font = normal_font
    ws.cell(row=start_row + 4, column=1, value="GSTIN:").font = bold_font
    ws.cell(row=start_row + 4, column=2, value=supplier.gst_number or "N/A").font = normal_font
    ws.cell(row=start_row + 5, column=1, value="Address:").font = bold_font
    ws.cell(row=start_row + 5, column=2, value=supplier.address or "N/A").font = normal_font

    headers = ["SKU", "Product Name", "Default Market Price (₹)", "Supplier Custom Cost (₹)", "Expiry Required"]
    row_idx = start_row + 7
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_idx == 5 else 'left')

    row_idx = start_row + 8
    for entry in catalog_items:
        ws.cell(row=row_idx, column=1, value=entry.item.sku).font = normal_font
        ws.cell(row=row_idx, column=2, value=entry.item.name).font = normal_font
        
        cell_mkt = ws.cell(row=row_idx, column=3, value=float(entry.item.unit_price))
        cell_mkt.font = normal_font
        cell_mkt.number_format = '"\u20B9"#,##0.00'
        
        cell_cost = ws.cell(row=row_idx, column=4, value=float(entry.price))
        cell_cost.font = bold_font
        cell_cost.number_format = '"\u20B9"#,##0.00'
        
        cell_exp = ws.cell(row=row_idx, column=5, value="Yes" if entry.item.has_expiry_date else "No")
        cell_exp.font = normal_font
        cell_exp.alignment = Alignment(horizontal='center')
        
        for col_idx in range(1, 6):
            ws.cell(row=row_idx, column=col_idx).border = thin_border
            
        row_idx += 1

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"Catalog_{supplier.name.replace(' ', '_')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

@login_required
@transaction.atomic
def auto_generate_pos(request):
    """
    Auto-generates Purchase Orders for all active products with a reorder alert.
    Orders the quantity matching the item's reorder threshold (minimum stock level).
    POs are prepared supplier-wise; deactivated/disabled suppliers are ignored.
    """
    if request.method != 'POST':
        return HttpResponseRedirect(reverse('inventory:dashboard'))

    # Find active products that have reached/fallen below their reorder threshold
    alert_items = InventoryItem.objects.filter(
        is_active=True,
        stock_level__lte=F('reorder_threshold')
    )

    # Find all catalog entries for these items from active suppliers
    catalog_entries = SupplierCatalogItem.objects.filter(
        item__in=alert_items,
        supplier__is_active=True
    ).select_related('supplier', 'item')

    # Group by supplier
    from collections import defaultdict
    supplier_groups = defaultdict(list)
    for entry in catalog_entries:
        supplier_groups[entry.supplier].append(entry)

    if not supplier_groups:
        messages.warning(request, "No low-stock items with active suppliers were found.")
        if request.headers.get('HX-Request'):
            response = HttpResponse()
            response['HX-Redirect'] = reverse('inventory:dashboard')
            return response
        return redirect('inventory:dashboard')

    # Get the initial draft state for Purchase Orders
    draft_state = State.objects.filter(
        workflow__model_name='inventory.PurchaseOrder',
        is_initial=True
    ).first()

    po_count = 0
    for supplier, entries in supplier_groups.items():
        # Create a new PurchaseOrder
        po = PurchaseOrder.objects.create(
            supplier=supplier,
            workflow_state=draft_state,
        )

        total_amount = 0.00
        for entry in entries:
            qty = entry.item.reorder_threshold
            price = entry.price

            POLineItem.objects.create(
                purchase_order=po,
                item=entry.item,
                quantity=qty,
                unit_price=price
            )
            total_amount += float(qty) * float(price)

        po.total_amount = total_amount
        po.save()
        po_count += 1

    messages.success(
        request,
        f"Successfully auto-generated {po_count} Purchase Order(s) for products with reorder alerts."
    )

    if request.headers.get('HX-Request'):
        response = HttpResponse()
        response['HX-Redirect'] = reverse('inventory:po_list')
        return response

    return redirect('inventory:po_list')




