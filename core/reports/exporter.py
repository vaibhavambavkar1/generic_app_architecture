import io
import csv
import json
from django.core.serializers.json import DjangoJSONEncoder
from django.db import models

class ReportExporter:
    """
    Handles report data rendering and formatting to HTML, PDF, Excel, CSV, and JSON.
    """
    @staticmethod
    def to_csv(data, headers):
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers.values())
        for row in data:
            writer.writerow([row.get(f, '') for f in headers.keys()])
        return io.BytesIO(buffer.getvalue().encode('utf-8'))

    @staticmethod
    def to_json(data):
        json_str = json.dumps(data, cls=DjangoJSONEncoder, indent=2)
        return io.BytesIO(json_str.encode('utf-8'))

    @staticmethod
    def to_excel(data, headers):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.append(list(headers.values()))
        # Bold styling
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for row in data:
            ws.append([row.get(f, '') for f in headers.keys()])
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def to_pdf(data, headers, title="Report"):
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        story = []
        styles = getSampleStyleSheet()
        
        story.append(Paragraph(title, styles['Title']))
        story.append(Spacer(1, 12))

        table_data = [list(headers.values())]
        for row in data:
            table_data.append([str(row.get(f, '')) for f in headers.keys()])

        t = Table(table_data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.grey),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
            ('BACKGROUND', (0,1), (-1,-1), colors.whitesmoke),
            ('GRID', (0,0), (-1,-1), 1, colors.lightgrey),
        ]))
        story.append(t)
        doc.build(story)
        buffer.seek(0)
        return buffer

    @staticmethod
    def to_html(data, headers, title="Report"):
        html_str = f"<html><head><title>{title}</title><style>table, th, td {{ border: 1px solid #ddd; border-collapse: collapse; padding: 8px; }} th {{ background-color: #f2f2f2; text-align: left; }}</style></head><body><h1>{title}</h1><table><thead><tr>"
        for label in headers.values():
            html_str += f"<th>{label}</th>"
        html_str += "</tr></thead><tbody>"
        for row in data:
            html_str += "<tr>"
            for f in headers.keys():
                html_str += f"<td>{row.get(f, '')}</td>"
            html_str += "</tr>"
        html_str += "</tbody></table></body></html>"
        return html_str


class DataExporter:
    """
    Generic Data Exporter for generating CSV, JSON, and Excel formats 
    from Django QuerySets or lists of objects (preserved for backward compatibility).
    """
    @staticmethod
    def _get_field_value(obj, field_path):
        if isinstance(obj, dict):
            return obj.get(field_path, "")
            
        val = obj
        for part in field_path.split('.'):
            if val is None:
                break
            val = getattr(val, part, None)
            if callable(val):
                val = val()
                
        if isinstance(val, models.Model):
            return str(val)
        return val

    @staticmethod
    def export_csv(queryset, fields):
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([label for _, label in fields])
        for obj in queryset:
            row = []
            for field, _ in fields:
                val = DataExporter._get_field_value(obj, field)
                row.append(str(val) if val is not None else '')
            writer.writerow(row)
        return io.BytesIO(buffer.getvalue().encode('utf-8'))
        
    @staticmethod
    def export_json(queryset, fields):
        data = []
        for obj in queryset:
            row = {}
            for field, label in fields:
                row[label] = DataExporter._get_field_value(obj, field)
            data.append(row)
        json_str = json.dumps(data, cls=DjangoJSONEncoder, indent=2)
        return io.BytesIO(json_str.encode('utf-8'))

    @staticmethod
    def export_excel(queryset, fields):
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Export Data"
        headers = [label for _, label in fields]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        import datetime
        for obj in queryset:
            row = []
            for field, _ in fields:
                val = DataExporter._get_field_value(obj, field)
                if isinstance(val, (dict, list, tuple)):
                    val = str(val)
                elif isinstance(val, datetime.datetime) and val.tzinfo is not None:
                    val = val.replace(tzinfo=None)
                row.append(val)
            sheet.append(row)
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer
